import os
import json
import math
import time
import logging
from pypdf import PdfReader
from google import genai
from dotenv import load_dotenv

load_dotenv()

api_key = os.getenv("GEMINI_API_KEY")
if not api_key:
    raise ValueError("GEMINI_API_KEY is missing in the backend/.env file!")

client = genai.Client(api_key=api_key)

# Uses uvicorn's logger so messages show up in the backend terminal.
logger = logging.getLogger("uvicorn.error")

# Models. The fallback is optional: if you set GEMINI_FALLBACK_MODEL in .env,
# it is tried automatically when the primary model keeps returning 503s.
PRIMARY_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.6-flash")
FALLBACK_MODEL = os.getenv("GEMINI_FALLBACK_MODEL")
EMBED_MODEL = "gemini-embedding-001"

BUSY_MESSAGE = "The AI is currently busy handling high demand. Please wait a few seconds and try again."
RETRY_DELAYS = (1.5, 3, 6)

indexed_store = {
    "chunks": [],
    "embeddings": [],
    "filename": None,
    "language": "english",
}

# Optional explicit paths for Windows (set in backend/.env):
#   TESSERACT_CMD=C:\Program Files\Tesseract-OCR\tesseract.exe
#   POPPLER_PATH=C:\...\poppler-xx\Library\bin
TESSERACT_CMD = os.getenv("TESSERACT_CMD")
POPPLER_PATH = os.getenv("POPPLER_PATH")

if TESSERACT_CMD:
    try:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD
    except ImportError:
        pass


# ---------------------------------------------------------------------------
# Gemini call helpers (retry + optional fallback model)
# ---------------------------------------------------------------------------

_TRANSIENT_MARKERS = (
    "503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED", "DEADLINE",
    "RemoteProtocolError", "ReadError", "ReadTimeout", "ConnectError", "ConnectTimeout",
    "incomplete", "Connection reset", "timed out",
)


def _is_transient(err: Exception) -> bool:
    s = f"{type(err).__name__}: {err}"
    return any(t in s for t in _TRANSIENT_MARKERS)


def call_gemini_with_retry(fn):
    """Retry helper for short calls such as embeddings."""
    last_error = None
    for attempt in range(len(RETRY_DELAYS) + 1):
        try:
            return fn()
        except Exception as e:
            last_error = e
            if _is_transient(e) and attempt < len(RETRY_DELAYS):
                time.sleep(RETRY_DELAYS[attempt])
                continue
            break
    if last_error is not None and not _is_transient(last_error):
        raise Exception(str(last_error))
    raise Exception(BUSY_MESSAGE)


def _models_to_try():
    return [PRIMARY_MODEL] + ([FALLBACK_MODEL] if FALLBACK_MODEL else [])


def generate_text(prompt: str) -> str:
    """Non-streaming generation with retries, then the optional fallback model."""
    last_error = None
    for model in _models_to_try():
        for attempt in range(len(RETRY_DELAYS) + 1):
            try:
                response = client.models.generate_content(model=model, contents=prompt)
                return response.text or ""
            except Exception as e:
                last_error = e
                if _is_transient(e) and attempt < len(RETRY_DELAYS):
                    time.sleep(RETRY_DELAYS[attempt])
                    continue
                break
    if last_error is not None and not _is_transient(last_error):
        raise Exception(str(last_error))
    raise Exception(BUSY_MESSAGE)


MAX_RESUMES = 2


def _continuation_prompt(prompt: str, so_far: str) -> str:
    return f"""{prompt}

[PARTIAL ANSWER ALREADY SHOWN TO THE STUDENT]:
{so_far}

The answer above was cut off by a connection problem. Continue it from exactly where it
stopped, in the same language and format. Do not repeat anything already written, and do
not add an introduction or any remark about the interruption."""


def stream_text(prompt: str):
    """Streaming generation with retries and automatic resume.

    - Errors before any text arrives are retried (then the optional fallback model).
    - If the stream breaks halfway, we ask the model to continue from where it
      stopped, so the student gets a complete answer instead of a cut-off one.
    """
    so_far = ""
    resumes = 0
    last_error = None

    for model in _models_to_try():
        attempt = 0
        while attempt <= len(RETRY_DELAYS):
            current_prompt = prompt if not so_far else _continuation_prompt(prompt, so_far)
            try:
                for chunk in client.models.generate_content_stream(model=model, contents=current_prompt):
                    if chunk.text:
                        so_far += chunk.text
                        yield chunk.text
                return
            except Exception as e:
                last_error = e
                logger.warning(
                    "Gemini stream error (model=%s, chars_sent=%d): %s: %s",
                    model, len(so_far), type(e).__name__, e,
                )
                if so_far:
                    resumes += 1
                    if resumes > MAX_RESUMES:
                        yield "\n\n_The connection was interrupted before the answer finished. Please ask again._"
                        return
                    time.sleep(1.0)
                    continue  # resume from where the text stopped
                if _is_transient(e) and attempt < len(RETRY_DELAYS):
                    time.sleep(RETRY_DELAYS[attempt])
                    attempt += 1
                    continue
                break  # give up on this model, try the fallback if any

    if last_error is not None and not _is_transient(last_error):
        yield str(last_error)
    else:
        yield BUSY_MESSAGE


def cosine_similarity(v1, v2):
    dot_product = sum(a * b for a, b in zip(v1, v2))
    magnitude1 = math.sqrt(sum(a * a for a in v1))
    magnitude2 = math.sqrt(sum(b * b for b in v2))
    if not magnitude1 or not magnitude2:
        return 0.0
    return dot_product / (magnitude1 * magnitude2)


# ---------------------------------------------------------------------------
# Language handling
# ---------------------------------------------------------------------------

def detect_document_language(text: str) -> str:
    """Returns "urdu" if the text is mostly Arabic-script (Urdu/Arabic),
    otherwise "english"."""
    sample = text[:6000]
    arabic = sum(1 for ch in sample if "\u0600" <= ch <= "\u06FF" or "\u0750" <= ch <= "\u077F")
    latin = sum(1 for ch in sample if ch.isascii() and ch.isalpha())
    total = arabic + latin
    if total == 0:
        return "english"
    return "urdu" if arabic / total > 0.4 else "english"


def normalize_language(language) -> str:
    language = (language or "auto").strip().lower()
    return language if language in ("auto", "english", "roman_urdu", "urdu") else "auto"


LANGUAGE_INSTRUCTIONS = {
    "english": "Write the entire response in clear, natural English.",
    "roman_urdu": (
        "Write the entire response in Roman Urdu (Urdu written with English/Latin "
        "letters, the way people write Urdu in messages in Pakistan)."
    ),
    "urdu": (
        "Write the entire response in Urdu script (اردو) using correct, natural Urdu. "
        "Do not answer in English."
    ),
}


def language_rule(language: str) -> str:
    language = normalize_language(language)
    if language in LANGUAGE_INSTRUCTIONS:
        return (
            f"- LANGUAGE (mandatory): {LANGUAGE_INSTRUCTIONS[language]} This applies even "
            "if the question text or the document is in a different language. Names, "
            "technical terms, and Arabic verses or duas may stay in their original script."
        )
    return (
        "- LANGUAGE: Detect the language/script of the STUDENT QUESTION and reply fully in "
        "that same language: Roman Urdu -> Roman Urdu, Urdu script -> Urdu script, "
        "English -> English. Never mix languages within one response."
    )


def response_guidelines(language: str) -> str:
    return f"""
[RESPONSE GUIDELINES]:
{language_rule(language)}
- Give a detailed, well-explained answer that teaches the concept, not just
  a terse fact. Favor depth and clarity, but avoid repetition or filler;
  every sentence should add information.
- If asked for a summary, give at least 8-10 well-organized bullet points
  covering the whole document, not only its beginning.
- If asked for MCQs, generate exactly the number requested, each with 4
  options, the correct answer clearly marked, and a one-line explanation.
- If asked for short-answer questions, generate exactly the number requested,
  each with a concise model answer.
- If asked to generate a QUIZ/PAPER, output ONLY the questions. Never
  include answers or solutions in that response.
- Any code, pseudocode, or syntax must be in a fenced Markdown code block
  with a language tag (e.g. ```python), never inline or plain text.
- Use bullet points, numbered lists, and bold for structure and key terms.
"""


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def ocr_pdf_page(page_image) -> str:
    """OCR one rendered page. Tries Urdu+English, falls back to English only
    if the Urdu language data is not installed."""
    import pytesseract

    try:
        return pytesseract.image_to_string(page_image, lang="urd+eng")
    except pytesseract.TesseractError:
        return pytesseract.image_to_string(page_image, lang="eng")


def extract_text_from_pdf(file_path: str):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    # Almost no text means a scanned/image PDF, so fall back to OCR.
    if len(text.strip()) < 20:
        try:
            from pdf2image import convert_from_path
        except ImportError:
            raise Exception(
                "This PDF is a scan with no embedded text, and OCR support "
                "(pdf2image/pytesseract) is not installed on the backend."
            )

        convert_kwargs = {"dpi": 300}
        if POPPLER_PATH:
            convert_kwargs["poppler_path"] = POPPLER_PATH

        try:
            images = convert_from_path(file_path, **convert_kwargs)
        except Exception as e:
            raise Exception(
                "Could not render this PDF for OCR. Check that Poppler is installed and "
                f"POPPLER_PATH is set in .env: {e}"
            )

        ocr_text = ""
        for img in images:
            page_text = ocr_pdf_page(img)
            if page_text:
                ocr_text += page_text + "\n"
        text = ocr_text

    if not text.strip():
        raise Exception(
            "Could not extract any text from this PDF, even with OCR. "
            "The scan quality may be too low."
        )

    return text, len(reader.pages)


def extract_text_from_docx(file_path: str):
    from docx import Document
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts), 1


def extract_text_from_xlsx(file_path: str):
    from openpyxl import load_workbook
    wb = load_workbook(file_path, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(cell) for cell in row if cell is not None)
            if row_text.strip():
                parts.append(row_text)
    return "\n".join(parts), len(wb.worksheets)


def extract_text_from_pptx(file_path: str):
    from pptx import Presentation
    prs = Presentation(file_path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
    return "\n".join(parts), len(prs.slides)


def process_document_and_index(file_path: str, filename: str):
    """Returns (chunk_count, unit_count, detected_language)."""
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        full_text, total_units = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        full_text, total_units = extract_text_from_docx(file_path)
    elif ext in (".xlsx", ".xls"):
        full_text, total_units = extract_text_from_xlsx(file_path)
    elif ext == ".pptx":
        full_text, total_units = extract_text_from_pptx(file_path)
    else:
        raise Exception(f"Unsupported file type: {ext}")

    if not full_text.strip():
        raise Exception(f"Could not extract any text from this {ext.upper()} file.")

    chunk_size = 800
    overlap = 100
    chunks = []
    for i in range(0, len(full_text), chunk_size - overlap):
        chunk = full_text[i:i + chunk_size]
        if chunk.strip():
            chunks.append(chunk)

    embeddings = []
    for chunk in chunks:
        response = call_gemini_with_retry(
            lambda c=chunk: client.models.embed_content(model=EMBED_MODEL, contents=c)
        )
        embeddings.append(response.embeddings[0].values)

    detected = detect_document_language(full_text)

    indexed_store["chunks"] = chunks
    indexed_store["embeddings"] = embeddings
    indexed_store["filename"] = filename
    indexed_store["language"] = detected

    return len(chunks), total_units, detected


# ---------------------------------------------------------------------------
# Retrieval
# ---------------------------------------------------------------------------

BROAD_KEYWORDS = (
    "summary", "summarize", "summarise", "overview", "key points", "quiz", "mcq",
    "exam question", "short-answer", "short answer", "questions from", "mind map",
    "خلاصہ", "khulasa",
)


def is_broad_request(question: str) -> bool:
    q = question.lower()
    return any(k in q for k in BROAD_KEYWORDS)


def _even_sample(chunks, max_chunks):
    n = len(chunks)
    if n <= max_chunks:
        return list(chunks)
    step = n / max_chunks
    return [chunks[int(i * step)] for i in range(max_chunks)]


def _get_context(question: str) -> str:
    if not indexed_store["chunks"]:
        raise Exception("Please upload and index a document first.")

    # Summaries, quizzes and mind maps need the whole document. Similarity
    # search on a word like "summary" returns random chunks, so for these we
    # sample evenly across the document instead.
    if is_broad_request(question):
        picked = _even_sample(indexed_store["chunks"], 12)
        return "\n\n--- Next Section ---\n\n".join(picked)

    top_k = min(5, len(indexed_store["chunks"]))
    q_response = call_gemini_with_retry(
        lambda: client.models.embed_content(model=EMBED_MODEL, contents=question)
    )
    q_embedding = q_response.embeddings[0].values

    similarities = []
    for idx, emb in enumerate(indexed_store["embeddings"]):
        similarities.append((cosine_similarity(q_embedding, emb), indexed_store["chunks"][idx]))
    similarities.sort(key=lambda x: x[0], reverse=True)
    return "\n\n--- Next Chunk ---\n\n".join(c for _, c in similarities[:top_k])


def build_prompt(question: str, context_text: str, language: str) -> str:
    return f"""
You are StudyMate AI, a smart personalized academic tutor for students.
Answer the question based strictly on the provided document context below.
If the context has OCR artifacts (odd spacing or misrecognized characters),
interpret the intended meaning instead of quoting the noise literally.

[DOCUMENT CONTEXT]:
{context_text}

[STUDENT QUESTION]:
{question}

{response_guidelines(language)}
"""


def answer_question_from_rag(question: str, language: str = "auto") -> str:
    context_text = _get_context(question)
    return generate_text(build_prompt(question, context_text, language))


def answer_question_from_rag_stream(question: str, language: str = "auto"):
    context_text = _get_context(question)
    yield from stream_text(build_prompt(question, context_text, language))


# ---------------------------------------------------------------------------
# Mind map
# ---------------------------------------------------------------------------

def generate_mind_map(language: str = "auto") -> dict:
    if not indexed_store["chunks"]:
        raise Exception("Please upload and index a document first.")

    context_text = "\n\n".join(_even_sample(indexed_store["chunks"], 12))
    language = normalize_language(language)
    if language in LANGUAGE_INSTRUCTIONS:
        label_language = LANGUAGE_INSTRUCTIONS[language].replace("the entire response", "every label")
    else:
        label_language = "Write every label in the same language as the document."

    prompt = f"""
You are StudyMate AI. Analyze the following document context and produce a concept mind map.

[DOCUMENT CONTEXT]:
{context_text}

Return ONLY valid JSON (no markdown fences, no extra commentary) in exactly this shape:
{{
  "central": "Main Topic Name",
  "branches": [
    {{"label": "Branch Concept", "children": ["Sub-point A", "Sub-point B"]}}
  ]
}}

Rules:
- 4 to 6 branches maximum.
- Each branch may have 0-3 children.
- Keep every label short (2-6 words).
- {label_language}
- Output must be parseable JSON and nothing else.
"""
    raw = generate_text(prompt).strip()

    start, end = raw.find("{"), raw.rfind("}")
    if start == -1 or end == -1:
        raise Exception("Could not generate a valid mind map. Please try again.")

    try:
        return json.loads(raw[start:end + 1])
    except json.JSONDecodeError:
        raise Exception("Could not generate a valid mind map. Please try again.")